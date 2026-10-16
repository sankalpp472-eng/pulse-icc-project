"""
Build Pulse_Workloads.xlsx: prompts, tool graphs and traces for workloads W1-W6.

  W1-W4: REAL agent trajectories shipped with the benchmarks
         tau-bench  historical_trajectories/  (GPT-4o x4, Claude 3.5 Sonnet x8 per task)
         tau2-bench data/tau2/results/final/  (GPT-4.1 and o4-mini, telecom, 4 trials)
  W5-W6: SYNTHETIC traces sampled from the graphs defined below (seeded).

Graph edges, probabilities and predictability metrics are Excel formulas over the
Trace_Steps sheet, so they update if trace rows are edited.

  python build_workloads_xlsx.py --tau-bench <path to tau-bench clone> \
                                 --tau2-bench <path to tau2-bench clone> [--out Pulse_Workloads.xlsx]

Both repos are MIT-licensed (Copyright Sierra).
"""

import argparse
import ast
import json
import random
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SEED = 252
N_SYNTH = 200
W1_WRITES = {"cancel_reservation", "update_reservation_flights",
             "update_reservation_passengers", "update_reservation_baggages"}

WORKLOADS = [
    # id, name, source, branch points, loops, expected predictability
    ("W1", "Airline: modify / cancel / service an existing reservation", "tau-bench airline (real traces)",
     "cabin class (basic economy cannot be modified); insurance (refund or not); flight already flown?; user accepts?",
     "repeated searches / reservation lookups", "Medium-high"),
    ("W2", "Airline: new booking", "tau-bench airline (real traces)",
     "direct flight found? else one-stop; payment type (certificate / gift card / card); free bags by membership tier",
     "repeated searches", "Medium"),
    ("W3", "Retail: returns, exchanges, order changes", "tau-bench retail (real traces)",
     "auth path (email vs name+zip); order pending vs delivered; return vs exchange vs modify; product variant lookups",
     "per-order / per-item lookups", "Medium-low"),
    ("W4", "Telecom troubleshooting", "tau2-bench telecom (real traces)",
     "each diagnostic result (status bar, APN, roaming, data usage, bills) decides the next check or fix",
     "check -> fix -> re-check (speed test, MMS test)", "Medium-low"),
    ("W5", "Research agent", "Synthetic (graph below, seeded sampling)",
     "page paywalled? -> new search; non-English? -> translate; enough evidence? -> search again or cite/write",
     "search <-> fetch <-> summarize loop", "Low-medium (repeat vs exit)"),
    ("W6", "Insurance claim / document pipeline", "Synthetic (graph below, seeded sampling)",
     "document type (3-way); valid?; fraud risk high -> human review; notify ends or triggers resubmission",
     "resubmission after invalid document", "Tunable (branch probabilities)"),
]

# ---------------------------------------------------------------- synthetic workloads

W5_GRAPH = {  # from -> [(to, probability, condition)]
    "START": [("web_search", 1.0, "always start with a search")],
    "web_search": [("fetch_page", 0.75, "promising HTML result"),
                   ("extract_pdf", 0.20, "top result is a PDF"),
                   ("web_search", 0.05, "no useful results -> reformulate query")],
    "fetch_page": [("summarize", 0.55, "readable English page"),
                   ("translate", 0.15, "page not in English"),
                   ("web_search", 0.20, "paywalled / irrelevant -> new search"),
                   ("fetch_page", 0.10, "follow a link on the page")],
    "extract_pdf": [("summarize", 0.70, "PDF parsed"),
                    ("translate", 0.10, "PDF not in English"),
                    ("web_search", 0.20, "scanned / unreadable PDF")],
    "translate": [("summarize", 1.0, "always summarize after translating")],
    "summarize": [("web_search", 0.45, "evidence insufficient -> search again"),
                  ("cite_lookup", 0.45, "enough evidence -> collect citations"),
                  ("fetch_page", 0.10, "open the next result of the same search")],
    "cite_lookup": [("cite_lookup", 0.30, "another citation needed"),
                    ("write_report", 0.70, "citations complete")],
    "write_report": [("END", 0.85, "report accepted"),
                     ("web_search", 0.15, "self-review finds a gap")],
}
W5_TOOLS = {
    "web_search": ("read", "Search the web; returns ranked result URLs with snippets."),
    "fetch_page": ("read", "Download an HTML page and return its main text (may be paywalled or non-English)."),
    "extract_pdf": ("read", "Download a PDF and extract its text (may fail on scanned PDFs)."),
    "translate": ("read", "Translate text to English."),
    "summarize": ("read", "Summarize gathered text and judge whether evidence is sufficient."),
    "cite_lookup": ("read", "Resolve a source to a formatted citation (DOI / publisher metadata)."),
    "write_report": ("write", "Write the final report with citations."),
}
W5_PROMPTS = [
    "Find recent peer-reviewed evidence on whether predictive pre-warming reduces serverless cold-start latency, and write a one-page summary with citations.",
    "Compare Knative and OpenFaaS for scale-to-zero support. Use official docs and at least two independent benchmarks; cite everything.",
    "What does current research say about LLM agent tool-call latency in production? Include at least one non-English source if relevant.",
    "Summarize the main findings of the Azure Functions 'Serverless in the Wild' paper and later work that builds on it, with citations.",
    "Write a short literature review on workflow graph prediction for LLM agents (2023 onward). Prefer PDFs of the original papers.",
    "Find government or industry reports on data-centre energy use from idle capacity, including European (German/French) sources, and summarize.",
    "Research how customer-support chatbots use tool calling (CRM, ticketing). Give 5 concrete examples with sources.",
    "Collect evidence for and against keeping serverless functions warm with periodic pings. Conclude with a recommendation and citations.",
    "Explain the Model Context Protocol (MCP) and how its servers are typically deployed, citing the specification and two deployment write-ups.",
    "Find benchmarks that measure multi-step agent reliability (e.g. tau-bench) and summarize what they measure, with links to the papers.",
]

W6_TOOLS = {
    "ocr": ("read", "Extract text and layout from an uploaded claim document (scan/photo/PDF)."),
    "classify_doc": ("read", "Classify the document: invoice, identity document, or other/unknown."),
    "validate_invoice": ("read", "Check invoice fields, totals and provider registry; returns valid/invalid."),
    "validate_id": ("read", "Check identity document number, expiry and name match; returns valid/invalid."),
    "fraud_score": ("read", "Score the claim for fraud risk (low/high) from history and document signals."),
    "human_review": ("write", "Route the claim to a human adjuster; returns approve/reject."),
    "approve_payout": ("write", "Approve the claim and schedule the payout."),
    "notify": ("write", "Notify the claimant (approval, rejection, or request to resubmit a document)."),
}
# W6 branch probabilities (the "dial"): change these and re-run to sweep predictability.
W6_P = {
    "doc_invoice": 0.50, "doc_id": 0.30, "doc_other": 0.20,   # classify_doc 3-way split
    "invoice_valid": 0.80, "id_valid": 0.85,                    # validation passes
    "fraud_low": 0.75,                                          # fraud_score -> approve directly
    "review_approve": 0.60,                                     # human_review -> approve
    "resubmit": 0.70,                                           # after invalid notice, claimant resubmits
}
W6_GRAPH_DOC = {  # for the Graph sheets' "condition" column
    ("START", "ocr"): "claim uploaded",
    ("ocr", "classify_doc"): "always",
    ("classify_doc", "validate_invoice"): f"invoice (p={W6_P['doc_invoice']})",
    ("classify_doc", "validate_id"): f"identity document (p={W6_P['doc_id']})",
    ("classify_doc", "human_review"): f"other / unknown type (p={W6_P['doc_other']})",
    ("validate_invoice", "fraud_score"): f"invoice valid (p={W6_P['invoice_valid']})",
    ("validate_invoice", "notify"): "invoice invalid -> ask for resubmission",
    ("validate_id", "fraud_score"): f"ID valid (p={W6_P['id_valid']})",
    ("validate_id", "notify"): "ID invalid -> ask for resubmission",
    ("fraud_score", "approve_payout"): f"low risk (p={W6_P['fraud_low']})",
    ("fraud_score", "human_review"): "high risk",
    ("human_review", "approve_payout"): f"adjuster approves (p={W6_P['review_approve']})",
    ("human_review", "notify"): "adjuster rejects",
    ("approve_payout", "notify"): "always notify approval",
    ("notify", "ocr"): f"HISTORY-DEPENDENT: after an invalid-document notice, claimant resubmits (p={W6_P['resubmit']})",
    ("notify", "END"): "after approval/rejection always; after invalid notice if claimant gives up",
}
W6_PROMPTS = [
    "Process claim CLM-1042: a scanned hospital invoice for INR 48,500 for an outpatient procedure.",
    "Process claim CLM-1043: the claimant uploaded a photo of their passport to complete KYC before payout.",
    "Process claim CLM-1044: a blurry phone photo of a pharmacy bill; amount looks handwritten.",
    "Process claim CLM-1045: a PDF that the claimant labelled 'discharge summary' (not an invoice).",
    "Process claim CLM-1046: a repeat claimant with three claims this month; invoice from a new clinic.",
    "Process claim CLM-1047: driving licence uploaded as ID; it expired last month.",
    "Process claim CLM-1048: dental invoice for INR 6,200 from a registered provider.",
    "Process claim CLM-1049: resubmitted invoice after an earlier rejection for a missing GST number.",
    "Process claim CLM-1050: an image of a handwritten receipt from a rural clinic.",
    "Process claim CLM-1051: a multi-page PDF bundle containing an invoice and an ID card.",
]


def sample_w5(rng: random.Random, max_steps: int = 30) -> list[str]:
    seq, node = [], "START"
    while True:
        if len(seq) >= max_steps - 2:          # cap long loops: wrap up the report
            return seq + (["cite_lookup"] if seq[-1:] != ["cite_lookup"] else []) + ["write_report"]
        outs = W5_GRAPH[node]
        node = rng.choices([o[0] for o in outs], [o[1] for o in outs])[0]
        if node == "END":
            return seq
        seq.append(node)


def sample_w6(rng: random.Random, max_rounds: int = 3) -> list[str]:
    seq = []
    for round_ in range(max_rounds):
        seq += ["ocr", "classify_doc"]
        r = rng.random()
        if r < W6_P["doc_invoice"]:
            seq.append("validate_invoice")
            valid = rng.random() < W6_P["invoice_valid"]
        elif r < W6_P["doc_invoice"] + W6_P["doc_id"]:
            seq.append("validate_id")
            valid = rng.random() < W6_P["id_valid"]
        else:                                   # unknown type -> straight to a human
            seq.append("human_review")
            if rng.random() < W6_P["review_approve"]:
                seq.append("approve_payout")
            return seq + ["notify"]
        if not valid:                           # invalid document -> ask to resubmit
            seq.append("notify")
            if round_ < max_rounds - 1 and rng.random() < W6_P["resubmit"]:
                continue
            return seq
        seq.append("fraud_score")
        if rng.random() < W6_P["fraud_low"]:
            seq.append("approve_payout")
        else:
            seq.append("human_review")
            if rng.random() < W6_P["review_approve"]:
                seq.append("approve_payout")
        return seq + ["notify"]
    return seq


# ---------------------------------------------------------------- real benchmark data

def tau_tool_defs(tools_dir: Path) -> dict:
    """name -> description, read from each tool's get_info() literal (no import needed)."""
    out = {}
    for f in sorted(tools_dir.glob("*.py")):
        tree = ast.parse(f.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if isinstance(node, ast.FunctionDef) and node.name == "get_info":
                ret = next(n for n in ast.walk(node) if isinstance(n, ast.Return))
                fn = ast.literal_eval(ret.value)["function"]
                out[fn["name"]] = fn["description"]
    return out


def tau2_tool_defs(path: Path) -> dict:
    """name -> (READ/WRITE/GENERIC, first docstring line) from @is_tool decorators."""
    txt = path.read_text(encoding="utf-8")
    pat = re.compile(r'@is_tool\(ToolType\.(\w+)\)\s*\n\s*def (\w+)\(.*?\).*?:\s*"""\s*(.*?)\n', re.S)
    return {m.group(2): (m.group(1), m.group(3).strip()) for m in pat.finditer(txt)}


def rw(name: str) -> str:
    return "read" if re.match(r"(get_|list_|search_|find_|calculate|think|check_|can_|run_)", name) else "write"


def load_tau(tau_root: Path):
    traces, prompts = [], {}
    for fname, model in [("gpt-4o", "gpt-4o"), ("sonnet-35-new", "claude-3.5-sonnet")]:
        for domain in ("airline", "retail"):
            data = json.loads((tau_root / "historical_trajectories" / f"{fname}-{domain}.json").read_text(encoding="utf-8"))
            for e in data:
                task = e["info"]["task"]
                gold = [a["name"] for a in task["actions"]]
                if domain == "retail":
                    wl = "W3"
                else:
                    wl = "W2" if "book_reservation" in gold and not W1_WRITES & set(gold) else "W1"
                seq = [tc["function"]["name"] for m in e["traj"] if m["role"] == "assistant"
                       for tc in (m.get("tool_calls") or [])]
                traces.append({"workload": wl, "source": f"tau-bench {domain} historical", "model": model,
                               "task_id": f"{domain}-{e['task_id']}", "trial": e["trial"],
                               "reward": e["reward"], "seq": seq})
                prompts.setdefault((wl, f"{domain}-{e['task_id']}"),
                                   {"text": task["instruction"], "gold": " > ".join(gold),
                                    "source": f"tau-bench {domain} (test split)", "scenario": ""})
    return traces, prompts


def load_tau2(tau2_root: Path):
    traces, prompts = [], {}
    res = tau2_root / "data" / "tau2" / "results" / "final"
    for f, model in [("gpt-4.1-2025-04-14", "gpt-4.1"), ("o4-mini-2025-04-16", "o4-mini")]:
        d = json.loads((res / f"{f}_telecom_op_gpt-4.1-2025-04-14_4trials.json").read_text(encoding="utf-8"))
        for t in d["tasks"]:
            ins = t["user_scenario"]["instructions"]
            text = " ".join(x for x in [ins.get("reason_for_call"), ins.get("known_info"),
                                        ins.get("task_instructions")] if x)
            gold = (t.get("evaluation_criteria") or {}).get("actions") or []
            m = re.match(r"\[(\w+)\]", t["id"])
            prompts.setdefault(("W4", t["id"]), {"text": text,
                                                "gold": " > ".join(a["name"] for a in gold),
                                                "source": "tau2-bench telecom", "scenario": m.group(1) if m else ""})
        for s in d["simulations"]:
            seq = [tc["name"] for msg in s["messages"] for tc in (msg.get("tool_calls") or [])]
            traces.append({"workload": "W4", "source": "tau2-bench telecom results", "model": model,
                           "task_id": s["task_id"], "trial": s["trial"],
                           "reward": (s.get("reward_info") or {}).get("reward"), "seq": seq})
    return traces, prompts


# ---------------------------------------------------------------- workbook

ARIAL = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=ARIAL, bold=True, color="FFFFFF")
BODY = Font(name=ARIAL, size=10)
INPUT_FONT = Font(name=ARIAL, size=10, color="0000FF")
TITLE = Font(name=ARIAL, bold=True, size=14)


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill = HDR_FONT, HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{row + 1}"  # a string: ws.cell() would create an empty row and shift appends


def style_body(ws, first_row, wrap_cols=()):
    for row in ws.iter_rows(min_row=first_row):
        for c in row:
            if c.font.color is None or c.font == Font():
                c.font = BODY
            if c.column in wrap_cols:
                c.alignment = Alignment(wrap_text=True, vertical="top")


def mermaid(edges: dict, min_p: float = 0.15) -> str:
    lines = ["flowchart LR"]
    for (a, b), p in sorted(edges.items(), key=lambda x: -x[1]):
        if p >= min_p:
            lines.append(f"  {a} -->|{p:.2f}| {b}")
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--tau-bench", required=True, type=Path)
    ap.add_argument("--tau2-bench", required=True, type=Path)
    ap.add_argument("--out", default=str(Path(__file__).with_name("Pulse_Workloads.xlsx")))
    a = ap.parse_args()

    traces, prompts = load_tau(a.tau_bench)
    t2, p2 = load_tau2(a.tau2_bench)
    traces += t2
    prompts.update(p2)

    rng = random.Random(SEED)
    for wl, sampler, plist in [("W5", sample_w5, W5_PROMPTS), ("W6", sample_w6, W6_PROMPTS)]:
        for i, text in enumerate(plist, 1):
            prompts[(wl, f"{wl}-P{i:02d}")] = {"text": text, "gold": "", "source": "synthetic (written for Pulse)",
                                               "scenario": ""}
        for i in range(N_SYNTH):
            traces.append({"workload": wl, "source": "synthetic (seeded sampling)", "model": "graph sampler",
                           "task_id": f"{wl}-P{(i % len(plist)) + 1:02d}", "trial": i // len(plist),
                           "reward": None, "seq": sampler(rng)})
    for i, t in enumerate(traces, 1):
        t["trace_id"] = f"{t['workload']}-T{i:05d}"

    # tool catalogue
    tools = []  # (workload, tool, actor, rw, description)
    air = tau_tool_defs(a.tau_bench / "tau_bench/envs/airline/tools")
    ret = tau_tool_defs(a.tau_bench / "tau_bench/envs/retail/tools")
    tel_agent = tau2_tool_defs(a.tau2_bench / "src/tau2/domains/telecom/tools.py")
    tel_user = tau2_tool_defs(a.tau2_bench / "src/tau2/domains/telecom/user_tools.py")
    for wl in ("W1", "W2"):
        tools += [(wl, n, "agent", rw(n), d) for n, d in air.items()]
    tools += [("W3", n, "agent", rw(n), d) for n, d in ret.items()]
    tools += [("W4", n, "agent", k.lower(), d) for n, (k, d) in tel_agent.items()]
    tools += [("W4", n, "user device", k.lower(), d) for n, (k, d) in tel_user.items()]
    tools += [("W5", n, "agent", k, d) for n, (k, d) in W5_TOOLS.items()]
    tools += [("W6", n, "agent", k, d) for n, (k, d) in W6_TOOLS.items()]

    # steps (START -> t1 -> ... -> END), edge list, python-side probabilities for mermaid text
    steps, edge_counts = [], defaultdict(Counter)
    for t in traces:
        path = ["START"] + t["seq"] + ["END"]
        for k in range(1, len(path)):
            steps.append((t["workload"], t["trace_id"], k, path[k - 1], path[k]))
            edge_counts[t["workload"]][(path[k - 1], path[k])] += 1

    wb = Workbook()

    # ---- README
    ws = wb.active
    ws.title = "README"
    rows = [
        ("Pulse workloads W1-W6: prompts, tool graphs and traces", TITLE),
        ("", None),
        ("Sheets", Font(name=ARIAL, bold=True)),
        ("Overview: one row per workload with branch points, source, measured predictability and a Mermaid graph.", None),
        ("Summary: per-workload metrics computed with formulas (traces, tool calls, success rate, entropy, top-1 accuracy).", None),
        ("Prompts: the user prompt for every task (real benchmark instructions for W1-W4, written prompts for W5-W6).", None),
        ("Tools: tool catalogue per workload (actor, read/write, description).", None),
        ("Traces: one row per agent run, with the full tool-call sequence.", None),
        ("Trace_Steps: one row per transition (from_tool -> to_tool), START/END included. Feed this to the predictor.", None),
        ("Edges: the workload graph as an edge list. count/probability/entropy are formulas over Trace_Steps.", None),
        ("G_W1 ... G_W6: the graph as a transition matrix P(next | current); rows = current tool, columns = next tool.", None),
        ("W6_Params: the branch probabilities used to generate W6 (the predictability dial).", None),
        ("", None),
        ("Sources", Font(name=ARIAL, bold=True)),
        ("W1/W2/W3: tau-bench historical_trajectories (GPT-4o x4 trials, Claude 3.5 Sonnet x8 trials per task), "
         "https://github.com/sierra-research/tau-bench (MIT, Copyright Sierra).", None),
        ("W4: tau2-bench data/tau2/results/final telecom runs (GPT-4.1 and o4-mini agents, GPT-4.1 user, 4 trials), "
         "https://github.com/sierra-research/tau2-bench (MIT, Copyright Sierra).", None),
        ("W5/W6: synthetic traces sampled from the graphs in this workbook (seed 252, 200 traces each).", None),
        ("W1 vs W2 split (airline): tasks whose gold actions contain book_reservation and no cancel/update are W2; all other airline tasks are W1. "
         "Only 4 of the 50 airline test tasks are pure new bookings, so W2 has just 48 traces: collect more with our own agent before relying on it.", None),
        ("W4 note: in tau2-bench telecom most tool calls are made by the CUSTOMER on their phone (actor = 'user device' in Tools). "
         "For Pulse, either deploy them as MCP tools too (device simulator) or filter to the agent's tools.", None),
        ("", None),
        ("How to rebuild", Font(name=ARIAL, bold=True)),
        ("python workloads/build_workloads_xlsx.py --tau-bench <tau-bench clone> --tau2-bench <tau2-bench clone>", None),
        ("To sweep W6 predictability, edit W6_P in the script and rebuild (W6_Params records the values used).", None),
    ]
    for i, (text, font) in enumerate(rows, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = font or BODY
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 140

    # ---- Trace_Steps (data first so other sheets can reference sizes)
    st = wb.create_sheet("Trace_Steps")
    header(st, 1, ["workload", "trace_id", "step", "from_tool", "to_tool"], [10, 14, 7, 30, 30])
    for r in steps:
        st.append(list(r))
    n_steps = len(steps) + 1
    S = lambda col: f"Trace_Steps!${col}$2:${col}${n_steps}"

    # ---- Traces
    tr = wb.create_sheet("Traces")
    header(tr, 1, ["trace_id", "workload", "source", "model", "task_id", "trial", "reward (1=solved)",
                   "n_tool_calls", "tool_sequence"], [14, 10, 28, 18, 34, 6, 10, 10, 120])
    for t in traces:
        tr.append([t["trace_id"], t["workload"], t["source"], t["model"], t["task_id"], t["trial"],
                   t["reward"], len(t["seq"]), " > ".join(t["seq"])])
    n_tr = len(traces) + 1
    T = lambda col: f"Traces!${col}$2:${col}${n_tr}"

    # ---- Edges (formulas)
    ed = wb.create_sheet("Edges")
    header(ed, 1, ["workload", "from_tool", "to_tool", "count", "from_total", "probability",
                   "entropy_term (bits*count)", "is_top_next", "condition / meaning"],
           [10, 30, 30, 9, 11, 12, 16, 11, 70])
    edge_rows = []
    for wl in [w[0] for w in WORKLOADS]:
        for (fr, to), _ in sorted(edge_counts[wl].items(), key=lambda x: (x[0][0] != "START", x[0][0], -x[1])):
            edge_rows.append((wl, fr, to))
    n_ed = len(edge_rows) + 1
    w5_cond = {(fr, to): c for fr, outs in W5_GRAPH.items() for to, _, c in outs}
    for i, (wl, fr, to) in enumerate(edge_rows, 2):
        cond = w5_cond.get((fr, to), "") if wl == "W5" else W6_GRAPH_DOC.get((fr, to), "") if wl == "W6" else ""
        ed.append([wl, fr, to,
                   f'=COUNTIFS({S("A")},A{i},{S("D")},B{i},{S("E")},C{i})',
                   f'=SUMIFS($D$2:$D${n_ed},$A$2:$A${n_ed},A{i},$B$2:$B${n_ed},B{i})',
                   f'=IF(E{i}=0,0,D{i}/E{i})',
                   f'=IF(F{i}=0,0,-D{i}*LOG(F{i},2))',
                   f'=IF(AND(D{i}=_xlfn.MAXIFS($D$2:$D${n_ed},$A$2:$A${n_ed},A{i},$B$2:$B${n_ed},B{i}),'
                   f'COUNTIFS($A$1:A{i - 1},A{i},$B$1:B{i - 1},B{i},$D$1:D{i - 1},D{i})=0),1,0)',
                   cond])
        ed.cell(row=i, column=6).number_format = "0.0%"
    E = lambda col: f"Edges!${col}$2:${col}${n_ed}"

    # ---- Summary (formulas)
    sm = wb.create_sheet("Summary", 1)
    header(sm, 1, ["workload", "name", "traces", "avg tool calls / trace", "success rate (real traces)",
                   "distinct tools used", "edges in graph", "H(next | current) bits",
                   "effective branching factor (2^H)", "top-1 next-tool accuracy (Markov-1)"],
           [10, 46, 9, 12, 14, 11, 10, 13, 15, 17])
    for i, (wl, name, *_rest) in enumerate(WORKLOADS, 2):
        sm.append([wl, name,
                   f'=COUNTIFS({T("B")},A{i})',
                   f'=AVERAGEIFS({T("H")},{T("B")},A{i})',
                   f'=IFERROR(AVERAGEIFS({T("G")},{T("B")},A{i}),"n/a (synthetic)")',
                   f'=SUMPRODUCT(({E("A")}=A{i})*({E("B")}<>"START")/COUNTIFS({E("A")},{E("A")},{E("B")},{E("B")}))',
                   f'=COUNTIFS({E("A")},A{i})',
                   f'=IFERROR(SUMIFS({E("G")},{E("A")},A{i})/SUMIFS({E("D")},{E("A")},A{i}),0)',
                   f'=2^H{i}',
                   f'=IFERROR(SUMIFS({E("D")},{E("A")},A{i},{E("H")},1)/SUMIFS({E("D")},{E("A")},A{i}),0)'])
        sm.cell(row=i, column=4).number_format = "0.0"
        sm.cell(row=i, column=5).number_format = "0.0%"
        sm.cell(row=i, column=8).number_format = "0.00"
        sm.cell(row=i, column=9).number_format = "0.0"
        sm.cell(row=i, column=10).number_format = "0.0%"
    note_row = len(WORKLOADS) + 3
    for k, text in enumerate([
        "How to read: H(next|current) is the conditional entropy of the next tool given the current one (0 = fully predictable).",
        "Effective branching factor 2^H ~ how many equally likely next tools there are. Top-1 accuracy = how often 'pick the most "
        "frequent next tool' is right: a simple baseline the Pulse predictor should beat.",
        "'distinct tools used' counts tools appearing as a from_tool. Success rate is the benchmark reward of the recorded agent runs.",
    ]):
        c = sm.cell(row=note_row + k, column=1, value=text)
        c.font = Font(name=ARIAL, size=9, italic=True)

    # ---- Overview
    ov = wb.create_sheet("Overview", 1)
    header(ov, 1, ["workload", "name", "source", "key branch points", "loops", "expected predictability",
                   "measured H(next|current) bits", "measured top-1 accuracy", "traces",
                   "Mermaid graph (edges with p >= 0.15; paste into mermaid.live)"],
           [10, 34, 26, 50, 26, 16, 14, 14, 9, 60])
    for i, (wl, name, src, br, loops, pred) in enumerate(WORKLOADS, 2):
        tot = defaultdict(int)
        for (fr, _), c in edge_counts[wl].items():
            tot[fr] += c
        probs = {(fr, to): c / tot[fr] for (fr, to), c in edge_counts[wl].items()}
        ov.append([wl, name, src, br, loops, pred, f"=Summary!H{i}", f"=Summary!J{i}", f"=Summary!C{i}",
                   mermaid(probs)])
        for col, fmt in [(7, "0.00"), (8, "0.0%")]:
            ov.cell(row=i, column=col).number_format = fmt
        for col in (7, 8, 9):
            ov.cell(row=i, column=col).font = Font(name=ARIAL, size=10, color="008000")
    ov.cell(row=len(WORKLOADS) + 3, column=1,
            value="Mermaid text is generated from the traces when the workbook is built (static). "
                  "Live numbers are in Summary / Edges / G_W* sheets.").font = Font(name=ARIAL, size=9, italic=True)

    # ---- Prompts
    pr = wb.create_sheet("Prompts")
    header(pr, 1, ["workload", "prompt_id / task_id", "source", "scenario", "user prompt (what the user asks)",
                   "gold actions (benchmark answer key)", "traces for this prompt"],
           [10, 34, 24, 18, 100, 60, 10])
    for i, ((wl, pid), p) in enumerate(sorted(prompts.items(), key=lambda x: (x[0][0], x[0][1])), 2):
        pr.append([wl, pid, p["source"], p["scenario"], p["text"], p["gold"],
                   f'=COUNTIFS({T("B")},A{i},{T("E")},B{i})'])

    # ---- Tools
    tl = wb.create_sheet("Tools")
    header(tl, 1, ["workload", "tool", "actor", "read / write", "description", "calls in traces"],
           [10, 32, 12, 10, 100, 12])
    for i, row in enumerate(tools, 2):
        tl.append(list(row) + [f'=COUNTIFS({S("A")},A{i},{S("E")},B{i})'])

    # ---- Graph matrices
    for wl in [w[0] for w in WORKLOADS]:
        g = wb.create_sheet(f"G_{wl}")
        froms = ["START"] + sorted({fr for fr, _ in edge_counts[wl] if fr != "START"},
                                   key=lambda x: -sum(c for (f, _), c in edge_counts[wl].items() if f == x))
        tos = [x for x in froms if x != "START"] + ["END"]
        tos += sorted({to for _, to in edge_counts[wl]} - set(tos))
        g.cell(row=1, column=1, value=f"{wl}: P(next tool | current tool), from Edges sheet. "
                                      f"Rows = current tool, columns = next tool.").font = Font(name=ARIAL, bold=True)
        g.cell(row=2, column=1, value="current \\ next").font = HDR_FONT
        g.cell(row=2, column=1).fill = HDR_FILL
        for j, to in enumerate(tos, 2):
            c = g.cell(row=2, column=j, value=to)
            c.font, c.fill = HDR_FONT, HDR_FILL
            c.alignment = Alignment(text_rotation=90, vertical="bottom", horizontal="center")
            g.column_dimensions[get_column_letter(j)].width = 6
        g.column_dimensions["A"].width = 32
        g.row_dimensions[2].height = 170
        for i, fr in enumerate(froms, 3):
            g.cell(row=i, column=1, value=fr).font = Font(name=ARIAL, bold=True, size=9)
            for j in range(2, len(tos) + 2):
                col = get_column_letter(j)
                c = g.cell(row=i, column=j,
                           value=f'=SUMIFS({E("F")},{E("A")},"{wl}",{E("B")},$A{i},{E("C")},{col}$2)')
                c.number_format = '0%;-0%;""'
                c.font = Font(name=ARIAL, size=8)
                c.alignment = Alignment(horizontal="center")
        last = f"{get_column_letter(len(tos) + 1)}{len(froms) + 2}"
        g.conditional_formatting.add(f"B3:{last}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                                  end_type="num", end_value=1, end_color="2E75B6"))
        g.freeze_panes = "B3"

    # ---- W6 params
    wp = wb.create_sheet("W6_Params")
    header(wp, 1, ["parameter", "value used", "meaning"], [18, 12, 80])
    meaning = {
        "doc_invoice": "classify_doc -> invoice", "doc_id": "classify_doc -> identity document",
        "doc_other": "classify_doc -> other/unknown (goes to human_review)",
        "invoice_valid": "validate_invoice passes", "id_valid": "validate_id passes",
        "fraud_low": "fraud_score low risk -> approve_payout directly", "review_approve": "human_review approves",
        "resubmit": "after an invalid-document notice, the claimant resubmits (loop back to ocr)",
    }
    for i, (k, v) in enumerate(W6_P.items(), 2):
        wp.append([k, v, meaning[k]])
        wp.cell(row=i, column=2).font = INPUT_FONT
        wp.cell(row=i, column=2).comment = Comment("Value used by build_workloads_xlsx.py (W6_P). "
                                                   "Edit in the script and rebuild to regenerate traces.", "Pulse")
    wp.cell(row=len(W6_P) + 3, column=1,
            value="Blue = generator input. Changing it here does NOT resample traces; edit W6_P in the script and rebuild."
            ).font = Font(name=ARIAL, size=9, italic=True)

    for wsx in wb.worksheets:
        style_body(wsx, 2, wrap_cols={"Prompts": (5, 6), "Overview": (2, 3, 4, 5, 10), "Tools": (5,)}.get(wsx.title, ()))
    order = ["README", "Overview", "Summary", "Prompts", "Tools", "Edges"] + \
            [f"G_{w[0]}" for w in WORKLOADS] + ["W6_Params", "Traces", "Trace_Steps"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    wb.save(a.out)
    print(f"wrote {a.out}: {len(traces)} traces, {len(steps)} steps, {len(edge_rows)} edges, {len(prompts)} prompts")
    for wl in [w[0] for w in WORKLOADS]:
        n = sum(1 for t in traces if t["workload"] == wl)
        print(f"  {wl}: {n} traces, {len(edge_counts[wl])} edges")


if __name__ == "__main__":
    main()
